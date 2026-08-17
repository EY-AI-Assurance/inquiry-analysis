from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Iterable

import openpyxl
import xlrd

from .common import (
    CorruptDocumentError,
    EmptyDocumentError,
    ParsedDocument,
    SourceChunk,
    format_tabular_row,
    normalize_text,
)


def _format_value(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return normalize_text(value).replace("\n", " ")


def _sheet_chunks(
    sheet_index: int,
    sheet_name: str,
    rows: Iterable[tuple[int, list[object]]],
    starting_order: int,
) -> list[SourceChunk]:
    chunks: list[SourceChunk] = []
    batch: list[str] = []
    batch_start = 1
    batch_end = 1
    order = starting_order
    headers: list[str] = []
    header_score = 0

    def flush() -> None:
        nonlocal batch, order
        if not batch:
            return
        order += 1
        chunks.append(
            SourceChunk(
                source_id=f"sheet-{sheet_index}-rows-{batch_start}-{batch_end}",
                locator=f"工作表“{sheet_name}”行 {batch_start}–{batch_end}",
                content="\n".join(batch),
                order=order,
            )
        )
        batch = []

    for row_number, values in rows:
        formatted = [_format_value(value) for value in values]
        if not any(formatted):
            continue
        line, headers, header_score = format_tabular_row(
            row_number, formatted, headers, header_score
        )
        if not batch:
            batch_start = row_number
        batch_end = row_number
        batch.append(line)
        if len(batch) >= 30:
            flush()
    flush()
    return chunks


def parse_xlsx(data: bytes) -> ParsedDocument:
    try:
        workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
        chunks: list[SourceChunk] = []
        for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
            rows = (
                (row_number, [cell.value for cell in row])
                for row_number, row in enumerate(sheet.iter_rows(), start=1)
            )
            chunks.extend(_sheet_chunks(sheet_index, sheet.title, rows, len(chunks)))
        workbook.close()
    except Exception as exc:
        raise CorruptDocumentError("Excel 文件无法读取，可能已损坏或格式不正确。") from exc

    if not chunks:
        raise EmptyDocumentError("Excel 文件中没有可分析的数据。")
    return ParsedDocument(chunks=chunks, warnings=[])


def parse_xls(data: bytes) -> ParsedDocument:
    try:
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        chunks: list[SourceChunk] = []
        for sheet_index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(sheet_index)
            rows = (
                (row_index + 1, sheet.row_values(row_index))
                for row_index in range(sheet.nrows)
            )
            chunks.extend(_sheet_chunks(sheet_index + 1, sheet.name, rows, len(chunks)))
        workbook.release_resources()
    except Exception as exc:
        raise CorruptDocumentError("旧版 Excel 文件无法读取，可能已损坏或格式不正确。") from exc

    if not chunks:
        raise EmptyDocumentError("Excel 文件中没有可分析的数据。")
    return ParsedDocument(chunks=chunks, warnings=[])
